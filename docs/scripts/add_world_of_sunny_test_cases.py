#!/usr/bin/env python3
"""Add World of Sunny (About) module test cases to Sunny Diamonds Test Cases workbook."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from world_of_sunny_test_cases_data import (
    PRECONDITIONS_LIST,
    TEST_CASES,
    TEST_DATA_NOTES,
)

SCRIPT_DIR = Path(__file__).resolve().parent
DOCS_DIR = SCRIPT_DIR.parent
WORKBOOK_PATH = DOCS_DIR / "Sunny Diamonds Test Cases .xlsx"
RESULTS_JSON = SCRIPT_DIR / "world-of-sunny-test-results.json"

SHEET_NAME = "World of Sunny"
ISSUES_SHEET_NAME = "World of Sunny Issues"
MODULE = "World of Sunny"
ROUTES = "/world-of-sunny (/about redirects here)"
DATA_SOURCE = "Strapi CMS (about-page single type)"

HEADER = [
    "Test Case ID",
    "Module",
    "Area / Feature",
    "Test Scenario",
    "Test Steps",
    "Expected Result",
    "Priority",
    "Type",
    "Preconditions",
    "Test Data",
    "Status",
    "Actual Result",
    "Notes",
]

COLUMN_WIDTHS = {
    "A": 14,
    "B": 16,
    "C": 22,
    "D": 42,
    "E": 52,
    "F": 48,
    "G": 10,
    "H": 12,
    "I": 28,
    "J": 24,
    "K": 10,
    "L": 28,
    "M": 32,
}

ISSUES_HEADER = [
    "Issue ID",
    "Jira Key",
    "Module",
    "Area / Feature",
    "Title",
    "Description",
    "Priority",
    "Status",
    "Test Case IDs",
    "Notes",
]

ISSUES_COLUMN_WIDTHS = {
    "A": 12,
    "B": 12,
    "C": 16,
    "D": 22,
    "E": 36,
    "F": 48,
    "G": 10,
    "H": 12,
    "I": 28,
    "J": 32,
}

KNOWN_GAPS = [
    "SD-1: Faces section hover over-zoom on desktop",
    "SD-2/SD-3: Since 1997 section description copy gaps",
    "SD-4: Handcrafted text card titles mismatch",
    "SD-5: Timeline default year should be 1997 (first milestone)",
    "SD-6: Hero animation not matching Figma",
    "SD-7: Faces mobile name/role visibility",
    "SD-8: Mobile second hero banner size",
    "SD-133: Mobile icon/text alignment in sections",
    "SD-141: Handcrafted divider vs Figma",
]

ISSUES = [
    ("WOS-ISSUE-001", "SD-1", MODULE, "F. Faces", "Faces hover over-zoom", "Desktop face cards zoom too aggressively on hover compared to Figma.", "P2", "Open", "WOS-053", ""),
    ("WOS-ISSUE-002", "SD-2", MODULE, "E. Since 1997", "Since 1997 description copy (1)", "First description block does not match approved CMS/Figma copy.", "P2", "Open", "WOS-041", ""),
    ("WOS-ISSUE-003", "SD-3", MODULE, "E. Since 1997", "Since 1997 description copy (2)", "Second description block does not match approved CMS/Figma copy.", "P2", "Open", "WOS-041", ""),
    ("WOS-ISSUE-004", "SD-4", MODULE, "G. Handcrafted", "Handcrafted card titles", "Text card titles in Handcrafted section do not match Figma.", "P2", "Open", "WOS-061", ""),
    ("WOS-ISSUE-005", "SD-5", MODULE, "H. Timeline", "Timeline default year", "Timeline should default to 1997 (first milestone); currently defaults to first CMS array item.", "P1", "Open", "WOS-071", ""),
    ("WOS-ISSUE-006", "SD-6", MODULE, "C. Hero", "Hero animation", "Hero entrance/animation does not match Figma motion design.", "P3", "Open", "WOS-025", ""),
    ("WOS-ISSUE-007", "SD-7", MODULE, "F. Faces", "Mobile name/role visibility", "On mobile, team member name and role are not consistently visible without hover.", "P1", "Open", "WOS-052", ""),
    ("WOS-ISSUE-008", "SD-8", MODULE, "E. Since 1997", "Mobile second hero banner", "Second hero banner size on mobile does not match Figma dimensions.", "P2", "Open", "WOS-044", ""),
    ("WOS-ISSUE-009", "SD-133", MODULE, "G. Handcrafted", "Mobile icon/text alignment", "Icons and text misaligned on mobile in handcrafted section.", "P2", "Open", "WOS-066", ""),
    ("WOS-ISSUE-010", "SD-141", MODULE, "G. Handcrafted", "Handcrafted divider", "Section divider styling does not match Figma.", "P3", "Open", "WOS-063", ""),
]


def load_status_overrides() -> dict[str, str]:
    if not RESULTS_JSON.exists():
        return {}
    try:
        data = json.loads(RESULTS_JSON.read_text(encoding="utf-8"))
        return {k: v.get("status", "") for k, v in data.items() if isinstance(v, dict)}
    except (json.JSONDecodeError, OSError):
        return {}


def style_header_row(ws, row: int, col_count: int) -> None:
    fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_section_title(ws, row: int, title: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADER))
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")


def write_summary_block(ws, start_row: int) -> int:
    row = start_row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADER))
    ws.cell(row=row, column=1, value=f"{MODULE} — Test Case Summary").font = Font(bold=True, size=16)
    row += 2

    for label, value in [
        ("Module", MODULE),
        ("Routes", ROUTES),
        ("Data Source", DATA_SOURCE),
        ("Total Test Cases", str(len(TEST_CASES))),
        ("Known Gaps / Jira", str(len(KNOWN_GAPS))),
    ]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Preconditions:").font = Font(bold=True)
    row += 1
    for pre_id, pre_desc in PRECONDITIONS_LIST:
        ws.cell(row=row, column=1, value=f"{pre_id}")
        ws.cell(row=row, column=2, value=pre_desc)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Test Data Notes:").font = Font(bold=True)
    row += 1
    for note in TEST_DATA_NOTES:
        ws.cell(row=row, column=1, value=f"• {note}")
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Known Gaps:").font = Font(bold=True)
    row += 1
    for gap in KNOWN_GAPS:
        ws.cell(row=row, column=1, value=f"• {gap}")
        row += 1
    return row + 1


def populate_test_cases_sheet(ws, status_overrides: dict[str, str]) -> None:
    if ws.max_row > 0:
        ws.delete_rows(1, ws.max_row)

    row = write_summary_block(ws, 1)
    style_section_title(ws, row, "TEST CASES")
    row += 1

    header_row = row
    for col, value in enumerate(HEADER, start=1):
        ws.cell(row=header_row, column=col, value=value)
    style_header_row(ws, header_row, len(HEADER))
    row += 1

    wrap = Alignment(vertical="top", wrap_text=True)
    for case in TEST_CASES:
        tc_id, area, scenario, steps, expected, priority, typ, preconds, test_data, notes = case
        status = status_overrides.get(tc_id, "")
        values = (
            tc_id, MODULE, area, scenario, steps, expected,
            priority, typ, preconds, test_data, status, "", notes,
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = wrap
        row += 1

    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = f"A{header_row + 1}"


def populate_issues_sheet(ws) -> None:
    if ws.max_row > 0:
        ws.delete_rows(1, ws.max_row)

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(ISSUES_HEADER))
    ws.cell(row=row, column=1, value=f"{MODULE} — Tracked Issues").font = Font(bold=True, size=16)
    row += 2

    header_row = row
    for col, value in enumerate(ISSUES_HEADER, start=1):
        ws.cell(row=header_row, column=col, value=value)
    style_header_row(ws, header_row, len(ISSUES_HEADER))
    row += 1

    wrap = Alignment(vertical="top", wrap_text=True)
    for issue in ISSUES:
        for col, value in enumerate(issue, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = wrap
        row += 1

    for col_letter, width in ISSUES_COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = f"A{header_row + 1}"


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    status_overrides = load_status_overrides()
    wb = load_workbook(WORKBOOK_PATH)

    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.create_sheet(SHEET_NAME)
    populate_test_cases_sheet(ws, status_overrides)

    if ISSUES_SHEET_NAME in wb.sheetnames:
        issues_ws = wb[ISSUES_SHEET_NAME]
    else:
        issues_ws = wb.create_sheet(ISSUES_SHEET_NAME)
    populate_issues_sheet(issues_ws)

    wb.save(WORKBOOK_PATH)
    print(f"Updated '{SHEET_NAME}' with {len(TEST_CASES)} test cases")
    print(f"Updated '{ISSUES_SHEET_NAME}' with {len(ISSUES)} issues")
    print(f"File: {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
