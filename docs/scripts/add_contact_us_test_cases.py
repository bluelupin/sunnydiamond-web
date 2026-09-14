#!/usr/bin/env python3
"""Add Contact Us test cases and issues sheets to Sunny Diamonds Test Cases workbook."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "Sunny Diamonds Test Cases .xlsx"
TEST_SHEET = "Contact Us"
ISSUES_SHEET = "Contact Us Issues"

MODULE = "Contact Us"
ROUTES = "/contact"
DATA_SOURCE = "Strapi CMS (contact-page, generic-forms) + BFF /api/generic-submissions/submit"
PRD = "docs/prd/PRD-contact-us-page.md"

PRECONDITIONS = [
    ("PRE-01", "Strapi CMS reachable with published contact-page single type"),
    ("PRE-02", "Generic form configured with matching formTag and reason dropdown options"),
    ("PRE-03", "Test on Desktop (>=1280px), Tablet (~768px), Mobile (<=390px)"),
    ("PRE-04", "Browsers: Chrome, Safari, Firefox, Edge (latest)"),
    ("PRE-05", "Guest and logged-in Magento customer accounts available for prefill tests"),
]

TEST_DATA = [
    ("Valid enquiry", "Name: Jane Doe, Phone: 9876543210 (+91), Email: jane@example.com, Reason: Product enquiry, Message: 10+ chars"),
    ("Invalid phone +91", "Phone: 12345 (too short) or 5123456789 (invalid leading digit)"),
    ("Invalid email", "test@, test@domain, missing@.com"),
    ("Long message", "501+ character message to test max length validation"),
    ("Logged-in user", "Customer with name, email, phone in Magento profile"),
    ("CMS inactive hero", "heroSection.isActive = false"),
    ("CMS inactive form", "formSection.isActive = false"),
]

HEADER = (
    "TC ID",
    "Module",
    "Feature Area",
    "Test Scenario",
    "Steps",
    "Expected Result",
    "Priority",
    "Type",
    "Preconditions",
    "Test Data",
    "Status",
    "Actual Result",
    "Notes",
)

# id, area, scenario, steps, expected, priority, type, preconds, test_data, notes
TEST_CASES = [
    # A. Page load & routing
    ("CONTACT-001", "A. Page Load", "Open Contact Us page", "Navigate to /contact", "Page loads with skeleton then CMS content; title 'Contact Us | Sunny Diamonds'", "P0", "Functional", "PRE-01", "—", ""),
    ("CONTACT-002", "A. Page Load", "ISR revalidation", "Check page.tsx revalidate setting", "revalidate = 300 (5 min ISR)", "P2", "Technical", "—", "—", ""),
    ("CONTACT-003", "A. Page Load", "Sitemap inclusion", "Open /sitemap.xml", "/contact URL listed", "P2", "SEO", "—", "—", ""),
    ("CONTACT-004", "A. Page Load", "Inbound navigation links", "Click Contact Us from footer, profile orders, error page", "All links route to /contact", "P1", "Navigation", "PRE-01", "—", ""),
    ("CONTACT-005", "A. Page Load", "CMS fetch failure", "Simulate Strapi outage for contact-page", "Page renders without sections; no crash; metadata uses site defaults", "P1", "Resilience", "PRE-01", "CMS inactive", "No user-facing error message"),
    ("CONTACT-006", "A. Page Load", "All sections disabled in CMS", "Disable hero, contact, form, visit in CMS", "Minimal empty page without console errors", "P1", "CMS", "PRE-01", "CMS inactive sections", ""),
    ("CONTACT-007", "A. Page Load", "Loading skeleton", "Hard refresh /contact on slow network", "ContactPageSkeleton shown with aria-busy=true", "P2", "UX", "PRE-01", "—", ""),

    # B. Hero section
    ("CONTACT-010", "B. Hero", "Hero renders from CMS", "Open /contact with hero active", "Hero h1 title visible; desktop/mobile image or video if configured", "P1", "CMS", "PRE-01", "—", ""),
    ("CONTACT-011", "B. Hero", "Hero hidden when inactive", "Set heroSection.isActive=false in CMS", "Hero section omitted; content shifts up", "P1", "CMS", "PRE-01", "CMS inactive hero", ""),
    ("CONTACT-012", "B. Hero", "Hero hidden without title", "Remove hero title in CMS", "Hero section not rendered", "P2", "CMS", "PRE-01", "—", ""),
    ("CONTACT-013", "B. Hero", "Hero LCP preload", "Inspect network on page load", "Hero images preloaded via preloadPlpHeroLcpImages", "P2", "Performance", "PRE-01", "—", ""),

    # C. Intro & contact cards
    ("CONTACT-020", "C. Contact Cards", "Intro text renders", "View intro below hero", "CMS introText displayed; mobile/desktop variants if different", "P1", "CMS", "PRE-01", "—", ""),
    ("CONTACT-021", "C. Contact Cards", "Call Us card", "Verify Call Us card title, hours, phone link", "Card shows phone number; hours parsed from Label: Value format", "P0", "Functional", "PRE-01", "—", ""),
    ("CONTACT-022", "C. Contact Cards", "Call Us desktop click", "On desktop (>=768px), click phone number", "Number copied to clipboard; toast 'Phone number copied'; no dialer", "P0", "Functional", "PRE-03", "—", "Verified in browser"),
    ("CONTACT-023", "C. Contact Cards", "Call Us mobile tap", "On mobile (<768px), tap phone number", "Opens native dialer via tel: link", "P0", "Functional", "PRE-03", "—", ""),
    ("CONTACT-024", "C. Contact Cards", "Viewport resize phone behavior", "Resize 767px ↔ 768px and interact with phone", "Behavior switches between copy and dial on re-render", "P1", "Responsive", "PRE-03", "—", ""),
    ("CONTACT-025", "C. Contact Cards", "Clipboard denied", "Block clipboard permission; click phone on desktop", "Toast 'Unable to copy' with manual copy guidance", "P2", "Edge Case", "PRE-03", "—", ""),
    ("CONTACT-026", "C. Contact Cards", "Email card mailto", "Click email CTA", "Opens mail client with mailto: link", "P0", "Functional", "PRE-01", "—", ""),
    ("CONTACT-027", "C. Contact Cards", "WhatsApp card", "Click WhatsApp CTA", "Opens wa.me or external WhatsApp URL in new tab", "P0", "Functional", "PRE-01", "—", ""),
    ("CONTACT-028", "C. Contact Cards", "Card grid layout desktop", "View cards at >=768px", "3-column grid; cards equal height", "P2", "UI", "PRE-03", "—", ""),
    ("CONTACT-029", "C. Contact Cards", "Card stack layout mobile", "View cards at <768px", "Stacked cards with dividers and icons", "P2", "UI", "PRE-03", "—", ""),
    ("CONTACT-030", "C. Contact Cards", "Inactive card omitted", "Set contactOption.isActive=false for one card", "Card not shown; remaining cards render", "P1", "CMS", "PRE-01", "—", ""),
    ("CONTACT-031", "C. Contact Cards", "Card without actionable value", "Card with empty/invalid value", "Card skipped by mapper", "P2", "CMS", "PRE-01", "—", ""),
    ("CONTACT-032", "C. Contact Cards", "Mobile title variant", "Check email card on mobile", "Uses card.title (mobileTitle from mapper not applied)", "P2", "UI Gap", "PRE-03", "—", "ISSUE: mobileTitle unused"),

    # D. Enquiry form — fields & UI
    ("CONTACT-040", "D. Form UI", "Form section renders", "Scroll to enquiry form", "Form title, all fields, consent, submit button visible", "P0", "Functional", "PRE-01, PRE-02", "—", ""),
    ("CONTACT-041", "D. Form UI", "Form hidden when inactive", "Disable formSection in CMS", "Form section not rendered", "P1", "CMS", "PRE-01", "CMS inactive form", ""),
    ("CONTACT-042", "D. Form UI", "Form hidden without formTag", "Remove formTag from CMS", "Form section omitted", "P1", "CMS", "PRE-01", "—", ""),
    ("CONTACT-043", "D. Form UI", "Field labels from CMS", "Inspect form labels", "Name, Phone, Email, Reason, Message labels match CMS dynamicFields", "P1", "CMS", "PRE-02", "—", ""),
    ("CONTACT-044", "D. Form UI", "Phone country code select", "Open country code dropdown", "Options: +91, +1, +44; aria-label='Country code'", "P1", "Functional", "PRE-01", "—", ""),
    ("CONTACT-045", "D. Form UI", "Reason dropdown options", "Open reason dropdown", "CMS options shown (e.g. Product enquiry, Order Support, Other)", "P0", "Functional", "PRE-02", "—", "Verified: 6 options in browser"),
    ("CONTACT-046", "D. Form UI", "Field placeholders", "Inspect phone and email placeholders", "Field-specific placeholders from CMS", "P1", "UI", "PRE-02", "—", "FAIL: both show 'Full Name'"),
    ("CONTACT-047", "D. Form UI", "Message placeholder", "Inspect message textarea placeholder", "CMS message placeholder shown", "P2", "UI", "PRE-02", "—", ""),
    ("CONTACT-048", "D. Form UI", "Desktop 2-column phone/email", "View form at >=768px", "Phone and email fields side by side", "P2", "Responsive", "PRE-03", "—", ""),
    ("CONTACT-049", "D. Form UI", "Submit button full width mobile", "View submit at <768px", "Submit button spans full width", "P2", "Responsive", "PRE-03", "—", ""),
    ("CONTACT-050", "D. Form UI", "Live form config refresh", "Update generic-forms reason options in CMS", "Client fetches updated options without republishing contact page", "P2", "CMS", "PRE-02", "—", ""),

    # E. Form validation
    ("CONTACT-060", "E. Validation", "Submit empty form", "Click Submit with all fields empty", "Inline validation errors on all required fields", "P0", "Validation", "PRE-01", "—", "FAIL: Submit disabled"),
    ("CONTACT-061", "E. Validation", "Name required", "Leave name empty; blur field", "Error shown after blur", "P0", "Validation", "PRE-01", "—", ""),
    ("CONTACT-062", "E. Validation", "Name min length", "Enter 1 char name; submit", "Error: min 2 characters", "P1", "Validation", "PRE-01", "A", ""),
    ("CONTACT-063", "E. Validation", "Name max length", "Enter 81+ char name", "Error or input limited to 80 chars", "P2", "Validation", "PRE-01", "—", ""),
    ("CONTACT-064", "E. Validation", "Name invalid chars", "Enter name with numbers/symbols", "Validation error for invalid characters", "P2", "Validation", "PRE-01", "John@123", ""),
    ("CONTACT-065", "E. Validation", "Phone required +91", "Leave phone empty; blur", "Phone required error", "P0", "Validation", "PRE-01", "—", ""),
    ("CONTACT-066", "E. Validation", "Phone +91 format", "Enter invalid +91 numbers", "Must be 10 digits starting 6-9", "P0", "Validation", "PRE-01", "Invalid phone +91", ""),
    ("CONTACT-067", "E. Validation", "Phone +1 format", "Switch to +1; enter invalid length", "Exactly 10 digits required", "P1", "Validation", "PRE-01", "—", ""),
    ("CONTACT-068", "E. Validation", "Phone +44 format", "Switch to +44; enter 9 digits", "10-11 digits required", "P1", "Validation", "PRE-01", "—", ""),
    ("CONTACT-069", "E. Validation", "Phone digit sanitization", "Type letters in phone field", "Only digits accepted; max length enforced per country", "P1", "Validation", "PRE-01", "—", ""),
    ("CONTACT-070", "E. Validation", "Email required", "Leave email empty; blur", "Email required error", "P0", "Validation", "PRE-01", "—", ""),
    ("CONTACT-071", "E. Validation", "Email format", "Enter malformed emails", "Invalid email error", "P0", "Validation", "PRE-01", "Invalid email", ""),
    ("CONTACT-072", "E. Validation", "Reason required", "Fill all fields except reason; enable submit; submit", "Error: Please select a reason", "P0", "Validation", "PRE-02", "—", ""),
    ("CONTACT-073", "E. Validation", "Message required", "Leave message empty; submit", "Message required error", "P0", "Validation", "PRE-01", "—", ""),
    ("CONTACT-074", "E. Validation", "Message min length", "Enter <10 char message", "Min 10 characters error", "P1", "Validation", "PRE-01", "short", ""),
    ("CONTACT-075", "E. Validation", "Message max length", "Enter 501+ char message and submit", "Max 500 characters error", "P1", "Validation", "PRE-01", "Long message", "No maxlength on textarea"),
    ("CONTACT-076", "E. Validation", "Consent required", "Fill form but leave consent unchecked; submit", "Error: Please accept the terms to continue", "P0", "Validation", "PRE-01", "—", ""),
    ("CONTACT-077", "E. Validation", "Errors on blur", "Tab through fields leaving invalid values", "Errors appear per field on blur before submit", "P1", "UX", "PRE-01", "—", ""),
    ("CONTACT-078", "E. Validation", "aria-invalid on errors", "Trigger field error", "aria-invalid=true and aria-describedby linked to error element", "P1", "Accessibility", "PRE-01", "—", ""),

    # F. Form submission
    ("CONTACT-080", "F. Submission", "Valid form submit", "Fill all required fields + consent; submit", "POST to /api/generic-submissions/submit; success toast; form resets", "P0", "Functional", "PRE-01, PRE-02", "Valid enquiry", ""),
    ("CONTACT-081", "F. Submission", "Submit payload shape", "Inspect network request on submit", "JSON: formTag, fullName, phone (+code), email, reasonForContact, message, consentAccepted, sourcePage", "P1", "API", "PRE-02", "Valid enquiry", ""),
    ("CONTACT-082", "F. Submission", "Submit loading state", "Submit valid form", "Submit button disabled during isSubmitting", "P1", "UX", "PRE-02", "Valid enquiry", ""),
    ("CONTACT-083", "F. Submission", "Submit network failure", "Block API; submit valid form", "Destructive toast 'Unable to send message' with error detail", "P0", "Resilience", "PRE-02", "Valid enquiry", ""),
    ("CONTACT-084", "F. Submission", "Success message from CMS", "Submit successfully", "Toast shows CMS successMessage as description", "P1", "CMS", "PRE-02", "Valid enquiry", "No toast title on success"),
    ("CONTACT-085", "F. Submission", "Form reset after success", "Submit successfully", "All fields cleared; consent unchecked", "P1", "Functional", "PRE-02", "Valid enquiry", "Country code resets to +91"),
    ("CONTACT-086", "F. Submission", "Double submit prevention", "Rapidly click submit twice", "Only one submission sent", "P1", "Functional", "PRE-02", "Valid enquiry", ""),

    # G. Consent & policies
    ("CONTACT-090", "G. Consent", "Consent checkbox renders", "View consent row", "Checkbox + CMS consent label with linked policies", "P0", "Functional", "PRE-01", "—", ""),
    ("CONTACT-091", "G. Consent", "Terms link navigation", "Click TERMS & CONDITIONS in consent", "Navigates to /terms-and-conditions", "P0", "Navigation", "PRE-01", "—", ""),
    ("CONTACT-092", "G. Consent", "Privacy link navigation", "Click PRIVACY POLICY in consent", "Navigates to policy hub privacy-policy route", "P0", "Navigation", "PRE-01", "—", ""),
    ("CONTACT-093", "G. Consent", "Checkbox matches cart style", "Compare checkbox with cart page", "Same GiftingPanelCheckbox component/style", "P2", "UI", "PRE-01", "—", ""),
    ("CONTACT-094", "G. Consent", "Consent without label dead end", "CMS: requiresConsent=true, consentLabel empty", "Consent UI hidden but submit permanently disabled", "P1", "Edge Case", "PRE-01", "—", "Known bug"),

    # H. Auth prefill
    ("CONTACT-100", "H. Prefill", "Logged-in prefill", "Sign in; open /contact", "Name, email, phone prefilled from Magento profile", "P0", "Functional", "PRE-05", "Logged-in user", ""),
    ("CONTACT-101", "H. Prefill", "Prefill does not overwrite typed data", "Type in name before prefill loads", "User-typed value preserved", "P1", "Functional", "PRE-05", "Logged-in user", ""),
    ("CONTACT-102", "H. Prefill", "Prefill once only", "Clear a prefilled field after load", "Field not re-filled from profile", "P2", "Functional", "PRE-05", "Logged-in user", ""),
    ("CONTACT-103", "H. Prefill", "Profile fetch failure", "Simulate profile API failure while logged in", "Form remains empty/editable; no blocking error", "P2", "Resilience", "PRE-05", "—", ""),

    # I. Visit Us
    ("CONTACT-110", "I. Visit Us", "Visit Us section renders", "Scroll to bottom of page", "Visit Us title, description, image, CTA visible", "P1", "Functional", "PRE-01", "—", ""),
    ("CONTACT-111", "I. Visit Us", "Visit Us hidden when inactive", "Disable visitSection in CMS", "Section omitted", "P1", "CMS", "PRE-01", "—", ""),
    ("CONTACT-112", "I. Visit Us", "Visit Us CTA", "Click Visit Us CTA", "Opens Book a Visit flow or store locator per CMS config", "P1", "Functional", "PRE-01", "—", ""),
    ("CONTACT-113", "I. Visit Us", "Shared PDP component", "Compare with PDP Visit Us", "Same ProductDetailVisitUsSection behavior", "P2", "Technical", "PRE-01", "—", ""),

    # J. SEO & metadata
    ("CONTACT-120", "J. SEO", "Page title from CMS", "View page source / meta", "CMS metaTitle in <title>", "P1", "SEO", "PRE-01", "—", "Verified: Contact Us | Sunny Diamonds"),
    ("CONTACT-121", "J. SEO", "Meta description from CMS", "Inspect meta description tag", "CMS metaDescription rendered", "P1", "SEO", "PRE-01", "—", ""),
    ("CONTACT-122", "J. SEO", "Canonical URL", "Inspect canonical link", "CMS canonical or /contact default", "P1", "SEO", "PRE-01", "—", ""),
    ("CONTACT-123", "J. SEO", "OG image", "Inspect og:image when configured", "CMS ogImage URL in metadata", "P2", "SEO", "PRE-01", "—", ""),
    ("CONTACT-124", "J. SEO", "Fallback metadata", "CMS SEO fields empty", "siteConfig defaults used", "P2", "SEO", "PRE-01", "—", ""),

    # K. Accessibility & mobile header
    ("CONTACT-130", "K. Accessibility", "Heading hierarchy", "Inspect page structure", "h1 hero, h2 form/cards/visit", "P1", "Accessibility", "PRE-01", "—", ""),
    ("CONTACT-131", "K. Accessibility", "Form field labels", "Screen reader / accessibility tree", "All inputs have associated labels", "P1", "Accessibility", "PRE-01", "—", ""),
    ("CONTACT-132", "K. Accessibility", "Mobile header account icon", "On mobile, tap user icon in header", "Should navigate to account/profile; aria-label matches destination", "P0", "Accessibility", "PRE-03", "—", "FAIL: links to /contact, label 'Account'"),
    ("CONTACT-133", "K. Accessibility", "Keyboard form navigation", "Tab through form fields and submit", "Logical tab order; focus visible", "P1", "Accessibility", "PRE-01", "—", ""),

    # L. CMS content quality
    ("CONTACT-140", "L. CMS Content", "Email CTA text accuracy", "Review Email Us card link text", "Properly formatted email address as link label", "P1", "Content", "PRE-01", "—", "FAIL: GET INTOUCH@SUNNTDIAMONDS.COM"),
    ("CONTACT-141", "L. CMS Content", "Spelling in card descriptions", "Review all card copy", "No spelling/grammar errors", "P2", "Content", "PRE-01", "—", "FAIL: 'assisstance' typo"),
    ("CONTACT-142", "L. CMS Content", "WhatsApp CTA label", "Review WhatsApp link text", "User-friendly label (not generic 'Whatsapp')", "P2", "Content", "PRE-01", "—", ""),

    # M. Security & architecture
    ("CONTACT-150", "M. Architecture", "BFF proxy for submit", "Inspect submit network call", "POST goes to /api/generic-submissions/submit (not direct Strapi)", "P1", "Security", "PRE-02", "—", ""),
    ("CONTACT-151", "M. Architecture", "Direct Strapi form fetch", "Inspect generic-forms client call", "Client calls Strapi directly for form refresh", "P2", "Architecture", "PRE-02", "—", "Inconsistent with BFF pattern"),
    ("CONTACT-152", "M. Architecture", "No server-side validation on BFF", "POST invalid payload to BFF", "BFF proxies to Strapi without validation", "P2", "Security", "—", "—", ""),
    ("CONTACT-153", "M. Architecture", "No rate limiting", "Rapid repeated submissions", "No client or BFF rate limiting", "P2", "Security", "PRE-02", "—", ""),

    # N. Known gaps
    ("CONTACT-160", "N. Known Gaps", "Automated tests", "Search codebase for contact tests", "No unit/e2e tests for contact module", "P2", "Gap", "—", "—", ""),
    ("CONTACT-161", "N. Known Gaps", "Reason label without options", "CMS reason label set, no dropdown options", "Form permanently unsubmittable", "P1", "Gap", "PRE-01", "—", ""),
    ("CONTACT-162", "N. Known Gaps", "Static content dead code", "Check features/contact/data/content.ts usage", "File not imported; CMS-only design confirmed", "P3", "Gap", "—", "—", ""),
]

ISSUES_HEADER = (
    "Issue ID",
    "Module",
    "Area",
    "Title",
    "Description",
    "Severity",
    "Type",
    "Status",
    "Related TCs",
    "Jira",
    "Notes",
)

ISSUES = [
    (
        "CONTACT-ISSUE-001",
        "Form",
        "Validation UX",
        "Submit button disabled when form is invalid",
        "Submit uses disabled={isSubmitting || !isFormReady}, preventing users from clicking Submit on an empty form. PRD test #9 expects inline validation on empty submit. Users only see errors field-by-field on blur.",
        "P0",
        "Bug",
        "Open",
        "CONTACT-060",
        "SD-162",
        "ContactFormSection.tsx line 418",
    ),
    (
        "CONTACT-ISSUE-002",
        "Form",
        "UI",
        "Phone and email placeholders show 'Full Name'",
        "ContactFormSection reuses form.fields.fieldPlaceholder (mapped from name field only) for phone and email inputs instead of field-specific placeholders.",
        "P1",
        "Bug",
        "Open",
        "CONTACT-046",
        "SD-163",
        "contact-page.mapper.ts maps only nameField.placeholder",
    ),
    (
        "CONTACT-ISSUE-003",
        "Navigation",
        "Accessibility",
        "Mobile header user icon mislabeled and wrong destination",
        "MobileHeaderBar links UserIcon to /contact with aria-label='Account'. Screen readers announce 'Account' but user lands on Contact Us.",
        "P0",
        "Bug",
        "Open",
        "CONTACT-132",
        "SD-164",
        "MobileHeaderBar.tsx line 52",
    ),
    (
        "CONTACT-ISSUE-004",
        "Form",
        "Consent",
        "Consent required but label missing creates dead end",
        "If requiresConsent=true and consentLabel is empty, checkbox UI is hidden but isFormReady still requires consentAccepted. Submit stays permanently disabled.",
        "P1",
        "Bug",
        "Open",
        "CONTACT-094",
        "SD-165",
        "PRD section 14 documents this edge case",
    ),
    (
        "CONTACT-ISSUE-005",
        "CMS Content",
        "Email card",
        "Malformed email CTA text",
        "Email Us card displays 'GET INTOUCH@SUNNTDIAMONDS.COM' instead of a properly formatted email address.",
        "P1",
        "Content",
        "Open",
        "CONTACT-140",
        "SD-166",
        "CMS configuration issue",
    ),
    (
        "CONTACT-ISSUE-006",
        "CMS Content",
        "Personal Concierge card",
        "Spelling error in card description",
        "Description reads 'Get quick assisstance from our dedicated member of our team' — typo and grammar issues.",
        "P2",
        "Content",
        "Open",
        "CONTACT-141",
        "SD-167",
        "CMS copy fix",
    ),
    (
        "CONTACT-ISSUE-007",
        "Contact Cards",
        "UI",
        "mobileTitle mapped but never rendered",
        "contact-page.mapper sets mobileTitle for email cards but ContactInfoSection always renders card.title.",
        "P2",
        "Gap",
        "Open",
        "CONTACT-032",
        "SD-168",
        "",
    ),
    (
        "CONTACT-ISSUE-008",
        "Architecture",
        "API",
        "Client-side direct Strapi call for generic-forms",
        "getGenericFormByTag calls Strapi from browser via apiFetch instead of BFF. Exposes CMS URL; depends on CORS.",
        "P2",
        "Architecture",
        "Open",
        "CONTACT-151",
        "SD-169",
        "",
    ),
    (
        "CONTACT-ISSUE-009",
        "Form",
        "Validation",
        "No maxlength on message textarea",
        "500-char limit validated only on submit; user can type beyond limit without immediate feedback.",
        "P2",
        "UX",
        "Open",
        "CONTACT-075",
        "SD-170",
        "",
    ),
    (
        "CONTACT-ISSUE-010",
        "Form",
        "API",
        "consentAccepted always sent as true",
        "Submit payload hardcodes consentAccepted: true regardless of CMS requiresConsent setting.",
        "P2",
        "Gap",
        "Open",
        "CONTACT-081",
        "SD-171",
        "",
    ),
    (
        "CONTACT-ISSUE-011",
        "Form",
        "Reset",
        "Country code resets to +91 after successful submit",
        "resetForm() always sets countryCode to +91; profile prefill flag resets so non-+91 users lose their code until remount.",
        "P2",
        "Bug",
        "Open",
        "CONTACT-085",
        "SD-172",
        "",
    ),
    (
        "CONTACT-ISSUE-012",
        "Form",
        "Toast",
        "Success toast has no title",
        "Success path only passes description to toast; error toast includes title 'Unable to send message'.",
        "P3",
        "UX",
        "Open",
        "CONTACT-084",
        "SD-173",
        "",
    ),
    (
        "CONTACT-ISSUE-013",
        "Contact Cards",
        "CSS",
        "Duplicate Tailwind classes in ContactInfoSection",
        "Section className includes md:px-10 twice (xl:px-[150px] md:px-10 md:px-10).",
        "P3",
        "Code Quality",
        "Open",
        "—",
        "SD-174",
        "ContactInfoSection.tsx line 20",
    ),
    (
        "CONTACT-ISSUE-014",
        "Testing",
        "Coverage",
        "No automated tests for Contact Us module",
        "No unit, integration, or e2e tests found for contact page, form, or submission.",
        "P2",
        "Gap",
        "Open",
        "CONTACT-160",
        "SD-175",
        "",
    ),
    (
        "CONTACT-ISSUE-015",
        "Resilience",
        "Error handling",
        "CMS failure shows empty page without user message",
        "getContactPage catch returns EMPTY_CONTACT_PAGE silently; no retry or error UI for users.",
        "P2",
        "UX",
        "Open",
        "CONTACT-005",
        "SD-176",
        "",
    ),
    (
        "CONTACT-ISSUE-016",
        "Form",
        "Validation",
        "Reason required with empty dropdown options",
        "If CMS sets reason label but no options, reasonRequired=true and form cannot be submitted.",
        "P1",
        "Bug",
        "Open",
        "CONTACT-161",
        "SD-177",
        "",
    ),
]

COLUMN_WIDTHS = {
    "A": 14,
    "B": 14,
    "C": 28,
    "D": 32,
    "E": 40,
    "F": 50,
    "G": 10,
    "H": 14,
    "I": 22,
    "J": 22,
    "K": 10,
    "L": 30,
    "M": 30,
}

ISSUES_COLUMN_WIDTHS = {
    "A": 18,
    "B": 12,
    "C": 14,
    "D": 36,
    "E": 48,
    "F": 10,
    "G": 12,
    "H": 10,
    "I": 14,
    "J": 12,
    "K": 30,
}


def style_header_row(ws, row: int, col_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F2937")
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def style_section_title(ws, row: int, title: str, merge_cols: int = 6) -> None:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_cols)


def write_test_cases_sheet(wb) -> None:
    if TEST_SHEET in wb.sheetnames:
        del wb[TEST_SHEET]
    ws = wb.create_sheet(TEST_SHEET)

    row = 1
    style_section_title(ws, row, "CONTACT US MODULE — TEST CASES")
    row += 1
    ws.cell(row=row, column=1, value="Module:")
    ws.cell(row=row, column=2, value=MODULE)
    ws.cell(row=row, column=3, value="Routes:")
    ws.cell(row=row, column=4, value=ROUTES)
    row += 1
    ws.cell(row=row, column=1, value="Data Source:")
    ws.cell(row=row, column=2, value=DATA_SOURCE)
    ws.cell(row=row, column=3, value="PRD:")
    ws.cell(row=row, column=4, value=PRD)
    row += 2

    style_section_title(ws, row, "GLOBAL PRECONDITIONS")
    row += 1
    ws.cell(row=row, column=1, value="ID")
    ws.cell(row=row, column=2, value="Description")
    style_header_row(ws, row, 2)
    row += 1
    for pre_id, desc in PRECONDITIONS:
        ws.cell(row=row, column=1, value=pre_id)
        ws.cell(row=row, column=2, value=desc)
        row += 1
    row += 1

    style_section_title(ws, row, "TEST DATA RECOMMENDATIONS")
    row += 1
    ws.cell(row=row, column=1, value="Dataset")
    ws.cell(row=row, column=2, value="Purpose")
    style_header_row(ws, row, 2)
    row += 1
    for dataset, purpose in TEST_DATA:
        ws.cell(row=row, column=1, value=dataset)
        ws.cell(row=row, column=2, value=purpose)
        row += 1
    row += 1

    style_section_title(ws, row, "TEST CASES")
    row += 1
    header_row = row
    for col, value in enumerate(HEADER, start=1):
        ws.cell(row=header_row, column=col, value=value)
    style_header_row(ws, header_row, len(HEADER))
    row += 1

    wrap = Alignment(vertical="top", wrap_text=True)
    for case in TEST_CASES:
        values = (*case, "", "")
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = wrap
        row += 1

    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = f"A{header_row + 1}"


def write_issues_sheet(wb) -> None:
    if ISSUES_SHEET in wb.sheetnames:
        del wb[ISSUES_SHEET]
    ws = wb.create_sheet(ISSUES_SHEET)

    row = 1
    style_section_title(ws, row, "CONTACT US MODULE — ISSUES / GAPS")
    row += 1
    ws.cell(row=row, column=1, value="Review Date:")
    ws.cell(row=row, column=2, value="2026-09-14")
    ws.cell(row=row, column=3, value="Total Issues:")
    ws.cell(row=row, column=4, value=len(ISSUES))
    row += 2

    style_section_title(ws, row, "ISSUE SUMMARY")
    row += 1
    p0 = sum(1 for i in ISSUES if i[5] == "P0")
    p1 = sum(1 for i in ISSUES if i[5] == "P1")
    p2 = sum(1 for i in ISSUES if i[5] == "P2")
    p3 = sum(1 for i in ISSUES if i[5] == "P3")
    for label, count in [("P0", p0), ("P1", p1), ("P2", p2), ("P3", p3)]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=count)
        row += 1
    row += 1

    style_section_title(ws, row, "ISSUES")
    row += 1
    header_row = row
    for col, value in enumerate(ISSUES_HEADER, start=1):
        ws.cell(row=header_row, column=col, value=value)
    style_header_row(ws, header_row, len(ISSUES_HEADER))
    row += 1

    wrap = Alignment(vertical="top", wrap_text=True)
    for issue in ISSUES:
        for col, value in enumerate(issue, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = wrap
        row += 1

    for col_letter, width in ISSUES_COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = f"A{header_row + 1}"


def main() -> None:
    wb = load_workbook(WORKBOOK_PATH)
    write_test_cases_sheet(wb)
    write_issues_sheet(wb)
    wb.save(WORKBOOK_PATH)
    print(f"Updated '{TEST_SHEET}' with {len(TEST_CASES)} test cases")
    print(f"Updated '{ISSUES_SHEET}' with {len(ISSUES)} issues")
    print(f"File: {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
